"""生成 RAG 评测前后对比表（xlsx）。

默认生成待填模板（30 条问句预填，命中列留空，汇总行为 Excel 公式）；
跑完评测后传入 results 自动填入 0/1：

    python -m eval.make_eval_xlsx \
        --results-a results/eval_baseline.jsonl --tag-a "baseline 512/64" \
        --results-b results/eval_tuned.jsonl    --tag-b "tuned 1024/128" \
        --out ../../docs/evidence/RAG评测对比表.xlsx
"""
import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).parent
HEADER_FILL = PatternFill("solid", start_color="1E40AF")
INPUT_FONT = Font(name="Arial", color="0000FF")  # 蓝色：待填/可改输入
BLACK_FONT = Font(name="Arial", color="000000")  # 黑色：公式与固定内容
THIN = Border(*[Side(style="thin", color="CBD5E1")] * 4)


def load_results(path: str | None) -> dict[str, dict]:
    if not path:
        return {}
    results = {}
    for line in (BASE_DIR / path).read_text(encoding="utf-8").splitlines() if not Path(path).is_absolute() else Path(path).read_text(encoding="utf-8").splitlines():
        if line.strip():
            r = json.loads(line)
            results[r["q"]] = r
    return results


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--results-a", help="A 组 eval jsonl（相对 eval/ 目录或绝对路径）")
    parser.add_argument("--results-b", help="B 组 eval jsonl")
    parser.add_argument("--tag-a", default="A组 chunk=512 overlap=64")
    parser.add_argument("--tag-b", default="B组 chunk=1024 overlap=128")
    parser.add_argument("--out", default=str(BASE_DIR / ".." / ".." / "docs" / "evidence" / "RAG评测对比表.xlsx"))
    args = parser.parse_args()

    res_a, res_b = load_results(args.results_a), load_results(args.results_b)
    questions = [
        json.loads(line)
        for line in (BASE_DIR / "questions.jsonl").read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "评测对比"

    ws["A1"] = "RAG 切分参数评测前后对比"
    ws["A1"].font = Font(name="Arial", size=14, bold=True)
    ws["A2"] = f"A组：{args.tag_a}"
    ws["A3"] = f"B组：{args.tag_b}"
    for r in (2, 3):
        ws.cell(row=r, column=1).font = INPUT_FONT

    headers = ["序号", "评测问句", "应召回文档", "A组召回命中", "A组回答命中", "B组召回命中", "B组回答命中"]
    ws.append([])
    ws.append(headers)
    header_row = 5
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row, column=col)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN

    first_data = header_row + 1
    for i, item in enumerate(questions, 1):
        row = first_data + i - 1
        ra, rb = res_a.get(item["q"]), res_b.get(item["q"])
        values = [
            i, item["q"], item["gold_doc"],
            int(ra["recall_hit"]) if ra else None,
            int(ra["answer_hit"]) if ra else None,
            int(rb["recall_hit"]) if rb else None,
            int(rb["answer_hit"]) if rb else None,
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = THIN
            cell.font = INPUT_FONT if col >= 4 else BLACK_FONT
            if col == 2:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    last_data = first_data + len(questions) - 1
    summary_row = last_data + 1
    ws.cell(row=summary_row, column=2, value="命中率（命中数/30）").font = Font(name="Arial", bold=True)
    for col in (4, 5, 6, 7):
        letter = get_column_letter(col)
        rng = f"{letter}{first_data}:{letter}{last_data}"
        cell = ws.cell(row=summary_row, column=col)
        cell.value = f'=IF(COUNT({rng})=0,"",SUM({rng})/{len(questions)})'
        cell.number_format = "0.0%"
        cell.font = BLACK_FONT
        cell.border = THIN

    widths = [6, 52, 16, 13, 13, 13, 13]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = f"A{first_data}"

    note = wb.create_sheet("指标说明")
    notes = [
        ["指标", "定义"],
        ["召回命中", "SSE citation 事件的来源文档集合包含应召回文档（gold_doc），填 1 否则 0"],
        ["回答命中", "模型回答文本命中任一 gold_keywords，填 1 否则 0"],
        ["评测方法", "backend/eval/run_eval.py，每题独立会话；两组切分参数需分别重建知识库后各跑一轮"],
        ["数据来源", "backend/eval/results/eval_{tag}.jsonl（跑完评测后用 make_eval_xlsx.py 自动填充）"],
    ]
    for row in notes:
        note.append(row)
    note.column_dimensions["A"].width = 14
    note.column_dimensions["B"].width = 90
    for cell in note[1]:
        cell.font = Font(name="Arial", bold=True)

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"written: {out}")


if __name__ == "__main__":
    main()
