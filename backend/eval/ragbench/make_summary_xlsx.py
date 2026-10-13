"""生成全迭代 7 指标汇总 xlsx（docs/evidence/RAG评估汇总表v2.xlsx）。

- 数据源：results/<tag>/summary.json（硬数值作为源数据）
- delta 列用 Excel 公式（=本轮-基线行），保证源数据更新后可重算
- 用法：python -m eval.ragbench.make_summary_xlsx
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RESULTS_DIR = Path(__file__).resolve().parents[1] / "results"
EVIDENCE_DIR = Path(__file__).resolve().parents[3] / "docs" / "evidence"

TAGS = [
    "baseline",
    "iter01-topk10", "iter02-nohybrid", "iter03-rerank", "iter04-rrf-k30",
    "iter05-chunk256", "iter06-chunk1024", "iter07-overlap128", "iter08-semantic",
    "iter09-auto-topk", "iter10-router", "iter11-context3", "iter12-prompt-concise",
    "iter13-router-f4", "iter14-router-f5", "best-combo",
]
METRICS = [
    ("recall_at_k", "召回率", "retrieval", "0.0%"),
    ("precision_at_k", "精准率", "retrieval", "0.0%"),
    ("mrr", "MRR", "retrieval", "0.000"),
    ("faithfulness", "忠实度", "generation", "0.0%"),
    ("answer_relevance", "答案相关性", "generation", "0.0%"),
    ("context_relevance", "上下文相关性", "generation", "0.0%"),
    ("context_recall", "上下文召回", "generation", "0.0%"),
]

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF")
DELTA_FILL = PatternFill("solid", fgColor="F2F2F2")


def main() -> None:
    runs = []
    for tag in TAGS:
        f = RESULTS_DIR / tag / "summary.json"
        if f.exists():
            runs.append(json.loads(f.read_text(encoding="utf-8")))
    if not runs or runs[0]["tag"] != "baseline":
        raise SystemExit("缺少基线结果")

    wb = Workbook()
    ws = wb.active
    ws.title = "全迭代汇总"

    n_m = len(METRICS)
    headers = (["实验", "变量变更"] + [label for _, label, _, _ in METRICS]
               + [f"Δ{label}" for _, label, _, _ in METRICS] + ["judged", "耗时(s)"])
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    for s in runs:
        row = [s["tag"], s.get("desc", "")]
        for key, _, layer, _ in METRICS:
            v = s.get(layer, {}).get(key)
            row.append(round(v, 4) if v is not None else "—")
        r = len(runs)  # placeholder, formulas added after append
        ws.append(row)
        excel_row = ws.max_row
        for i, (key, _, _, _) in enumerate(METRICS):
            col = get_column_letter(3 + i)
            delta_col = get_column_letter(3 + n_m + i)
            if excel_row == 2:
                ws[f"{delta_col}{excel_row}"] = "—"
            else:
                ws[f"{delta_col}{excel_row}"] = (
                    f'=IF(ISNUMBER({col}{excel_row}),{col}{excel_row}-{col}$2,"—")'
                )
        ws.cell(row=excel_row, column=3 + 2 * n_m + 1, value=s.get("generation", {}).get("judged", 0))
        ws.cell(row=excel_row, column=3 + 2 * n_m + 2, value=s.get("elapsed_s"))

    # 数字格式与字体
    for i, (_, _, _, fmt) in enumerate(METRICS):
        for row in ws.iter_rows(min_row=2, min_col=3 + i, max_col=3 + i, max_row=ws.max_row):
            for cell in row:
                cell.number_format = fmt
                cell.font = Font(name=FONT)
    for i in range(n_m):
        for row in ws.iter_rows(min_row=2, min_col=3 + n_m + i, max_col=3 + n_m + i, max_row=ws.max_row):
            for cell in row:
                cell.number_format = METRICS[i][3].replace("0.0%", "+0.0%;-0.0%;0.0%").replace("0.000", "+0.000;-0.000;0.000")
                cell.font = Font(name=FONT)
                cell.fill = DELTA_FILL
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=2, max_row=ws.max_row):
        for cell in row:
            cell.font = Font(name=FONT)

    # 各指标列最优值加粗
    for i in range(n_m):
        col = 3 + i
        vals = [(r, ws.cell(row=r, column=col).value)
                for r in range(2, ws.max_row + 1)
                if isinstance(ws.cell(row=r, column=col).value, (int, float))]
        if vals:
            best = max(vals, key=lambda x: x[1])[0]
            ws.cell(row=best, column=col).font = Font(name=FONT, bold=True)

    widths = [22, 46] + [11] * n_m + [11] * n_m + [8, 9]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"

    # 说明 sheet
    note = wb.create_sheet("说明")
    notes = [
        ["数据源", "backend/eval/results/<tag>/summary.json（评估集 questions_v2.jsonl，50 题）"],
        ["评估日期", runs[0].get("finished_at", "")],
        ["基线定义", "系统默认配置：hybrid+BM25、无 rerank、top_k=5 固定、chunk 512/64 window"],
        ["iter03 说明", "reranker（BAAI/bge-reranker-base 1.06GB ONNX）本机加载失败（bad allocation），"
                       "按系统设计降级为粗排，该轮仅检索层数据，不代表 rerank 真实效果"],
        ["judge 校准", "见 docs/evidence/rag-v2-judge校准记录.md（答案相关性 prompt 已按校准修订）"],
        ["忠实度已知行为", "从严取向：逻辑蕴含但未明确陈述的论断可能判不支持"],
    ]
    for r, (k, v) in enumerate(notes, 1):
        note.cell(row=r, column=1, value=k).font = Font(name=FONT, bold=True)
        note.cell(row=r, column=2, value=v).font = Font(name=FONT)
    note.column_dimensions["A"].width = 16
    note.column_dimensions["B"].width = 110

    EVIDENCE_DIR.mkdir(parents=True, exist_ok=True)
    out = EVIDENCE_DIR / "RAG评估汇总表v2.xlsx"
    wb.save(out)
    print(f"已生成 {out}（{len(runs)} 轮实验）")


if __name__ == "__main__":
    main()
